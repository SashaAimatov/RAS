import json
import io
import calendar
import logging
import time
import os
from datetime import datetime, date, timedelta
from pathlib import Path
from urllib.parse import urlencode

import requests
from docx import Document
from aiogram import Bot, Dispatcher, BaseMiddleware
from aiogram.filters import CommandStart, Command
from aiogram.types import (
    Message, InlineKeyboardMarkup, InlineKeyboardButton,
    CallbackQuery, InlineQuery, InlineQueryResultArticle,
    InputTextMessageContent
)
from apscheduler.schedulers.asyncio import AsyncIOScheduler

# ================= НАСТРОЙКИ =================
TELEGRAM_TOKEN = os.environ.get("TELEGRAM_TOKEN", "")
WEBHOOK_SECRET = os.environ.get("WEBHOOK_SECRET", "secret")

OWNER_USERNAME = "pampampa1"
INITIAL_ADMINS = ["ginytimary"]
CONTACT_USERNAME = "raspisach"

SUBSCRIBERS_FILE = Path("subscribers.json")
ADMINS_FILE = Path("admins.json")
BANNED_FILE = Path("banned.json")
BANNED_USERNAMES_FILE = Path("banned_usernames.json")
CANCELLED_FILE = Path("cancelled.json")
MUTE_FILE = Path("mute.json")
SETTINGS_FILE = Path("settings.json")
OWNER_FILE = Path("owner.json")

DEFAULT_SETTINGS = {
    "ref_monday": "2026-10-12",
    "ref_is_numerator": True,
}

WEEKDAYS_RU = ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье"]
WEEKDAYS_ALIASES = {
    "пн": "Понедельник", "понедельник": "Понедельник",
    "вт": "Вторник", "вторник": "Вторник",
    "ср": "Среда", "среда": "Среда",
    "чт": "Четверг", "четверг": "Четверг",
    "пт": "Пятница", "пятница": "Пятница",
}

THIRD_BUILDING_DAYS = {"Понедельник", "Пятница"}
YADISK_PUBLIC_KEY = "https://disk.yandex.ru/d/vCkixjXynkqHwQ"
MONTHS_RU_SHORT = ["Янв", "Фев", "Мар", "Апр", "Май", "Июн",
                   "Июл", "Авг", "Сен", "Окт", "Ноя", "Дек"]
START_TIME = time.time()

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler()],
)
log = logging.getLogger("bot")

# ================= РАСПИСАНИЕ =================
SCHEDULE = {
    "group": "ТМ-25-202",
    "times": {
        "1": "8:20–9:50", "2": "10:00–11:30", "3": "12:10–13:40",
        "4": "14:00–15:30", "5": "15:40–17:10", "6": "17:20–18:50",
        "7": "19:00–20:30"
    },
    "days": {
        "Понедельник": [
            {"pair": 2, "subject": "История России", "room": "107/3к/у"},
            {"pair": 3, "subject": "Процессы формообразования и инструменты", "room": "201м"},
            {"pair": 4, "subject": "Охрана труда", "room": "207м"},
            {"pair": 5, "subject": "Инженерная графика", "room": "207м"}
        ],
        "Вторник": [
            {"pair": 1, "subject": "Иностранный язык", "room": "120о / 03"},
            {"pair": 2,
             "numerator":   {"subject": "Математика в профессиональной деятельности", "room": "214"},
             "denominator": {"subject": "Технология машиностроения", "room": "122о"}},
            {"pair": 3, "subject": "Математика в профессиональной деятельности", "room": "214"},
            {"pair": 4, "subject": "Физическая культура", "room": ""}
        ],
        "Среда": [
            {"pair": 3, "subject": "Техническая механика", "room": "106о"},
            {"pair": 4,
             "numerator":   {"subject": "Техническая механика", "room": "201м"},
             "denominator": {"subject": "Инженерная графика", "room": "105о"}}
        ],
        "Четверг": [
            {"pair": 1,
             "numerator":   {"subject": "История России", "room": "310"},
             "denominator": {"subject": "—", "room": ""}},
            {"pair": 2, "subject": "Материаловедение", "room": "01"},
            {"pair": 3, "subject": "Техническая механика", "room": "106о"},
            {"pair": 4, "subject": "Материаловедение", "room": "01"},
            {"pair": 5,
             "numerator":   {"subject": "—", "room": ""},
             "denominator": {"subject": "Материаловедение", "room": "01"}}
        ],
        "Пятница": [
            {"pair": 3, "subject": "Технологическое оборудование", "room": "201м"},
            {"pair": 4, "subject": "Технология машиностроения", "room": "201м"},
            {"pair": 5, "subject": "Процессы формообразования и инструменты", "room": "201м"},
            {"pair": 6, "subject": "Основы бережливого производства", "room": "201м"}
        ]
    }
}

TIMES = SCHEDULE["times"]
DAYS = SCHEDULE["days"]

# ================= ФАЙЛЫ =================
def _load_json(path, default):
    if not Path(path).exists():
        return default
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return default


def _save_json(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def load_subscribers(): return set(_load_json(SUBSCRIBERS_FILE, []))
def save_subscribers(s): _save_json(SUBSCRIBERS_FILE, sorted(s))
def load_admins(): return set(_load_json(ADMINS_FILE, INITIAL_ADMINS))
def save_admins(s): _save_json(ADMINS_FILE, sorted(s))
def load_banned(): return set(_load_json(BANNED_FILE, []))
def save_banned(s): _save_json(BANNED_FILE, sorted(s))
def load_banned_usernames(): return set(_load_json(BANNED_USERNAMES_FILE, []))
def save_banned_usernames(s): _save_json(BANNED_USERNAMES_FILE, sorted(s))
def load_cancelled(): return _load_json(CANCELLED_FILE, {})
def save_cancelled(d): _save_json(CANCELLED_FILE, d)
def load_mute(): return _load_json(MUTE_FILE, {})
def save_mute(d): _save_json(MUTE_FILE, d)


def load_settings():
    return {**DEFAULT_SETTINGS, **_load_json(SETTINGS_FILE, {})}


def save_settings(s): _save_json(SETTINGS_FILE, s)


def add_subscriber(chat_id):
    subs = load_subscribers()
    if chat_id not in subs:
        subs.add(chat_id); save_subscribers(subs)


def remove_subscriber(chat_id):
    subs = load_subscribers()
    if chat_id in subs:
        subs.remove(chat_id); save_subscribers(subs)


def get_owner_chat_id(): return _load_json(OWNER_FILE, {}).get("chat_id")
def set_owner_chat_id(cid): _save_json(OWNER_FILE, {"chat_id": cid})
def _username(m): return (m.from_user.username or "").lower()
def is_owner(m): return _username(m) == OWNER_USERNAME.lower()
def is_admin(m): return is_owner(m) or _username(m) in load_admins()


def is_banned(m):
    if m.chat.id in load_banned(): return True
    u = _username(m)
    return bool(u and u in load_banned_usernames())


def is_muted(chat_id):
    until = load_mute().get(str(chat_id))
    if not until: return False
    try:
        return datetime.strptime(until, "%Y-%m-%d").date() >= date.today()
    except Exception:
        return False


# ================= НЕДЕЛЯ =================
def get_ref_monday():
    try:
        return datetime.strptime(load_settings()["ref_monday"], "%Y-%m-%d").date()
    except Exception:
        return date(2026, 10, 12)


def ref_is_numerator(): return load_settings().get("ref_is_numerator", True)


def is_numerator(d):
    weeks = (d - get_ref_monday()).days // 7
    return (weeks % 2 == 0) == ref_is_numerator()


def days_until_week_change(d):
    x = (7 - d.weekday()) % 7
    return x if x > 0 else 7


# ================= ЯНДЕКС.ДИСК =================
def yadisk_list_items(public_key, path="/"):
    params = {"public_key": public_key, "limit": 200}
    if path and path != "/": params["path"] = path
    url = "https://cloud-api.yandex.net/v1/disk/public/resources?" + urlencode(params)
    try:
        r = requests.get(url, timeout=20)
        if r.status_code != 200: return []
        return r.json().get("_embedded", {}).get("items", [])
    except Exception as e:
        log.exception(f"Яндекс.Диск список: {e}")
        return []


def yadisk_download(public_key, path):
    params = {"public_key": public_key, "path": path}
    url = "https://cloud-api.yandex.net/v1/disk/public/resources/download?" + urlencode(params)
    try:
        r = requests.get(url, timeout=20)
        if r.status_code != 200: return None
        href = r.json().get("href")
        if not href: return None
        fr = requests.get(href, timeout=40)
        return fr.content if fr.status_code == 200 else None
    except Exception as e:
        log.exception(f"Яндекс.Диск download: {e}")
        return None


MONTHS_RU_GEN = {1:"января",2:"февраля",3:"марта",4:"апреля",5:"мая",6:"июня",
                 7:"июля",8:"августа",9:"сентября",10:"октября",11:"ноября",12:"декабря"}


def find_replacement_file(target_date):
    day = target_date.day
    mn = MONTHS_RU_GEN[target_date.month]
    mf = mn[:-1] + "ь"
    wd = WEEKDAYS_RU[target_date.weekday()]
    is3 = wd in THIRD_BUILDING_DAYS
    cands = [f"{day} {mn}", f"{day}.{target_date.month:02d}.{target_date.year}", f"{day}.{target_date.month}.{target_date.year}"]
    has_date = lambda n: any(c in n for c in cands)
    is_docx = lambda n: n.lower().endswith(".docx")
    is_xlsx = lambda n: n.lower().endswith(".xlsx")
    is3k = lambda n: "3 корпус" in n.lower()
    files = []
    root = yadisk_list_items(YADISK_PUBLIC_KEY, "/")
    for it in root:
        if it.get("type") == "file": files.append(it)
    mp = None
    for it in root:
        if it.get("type") == "dir" and it.get("name", "").lower() in (mf.lower(), mn.lower()):
            mp = it.get("path"); break
    if mp:
        for it in yadisk_list_items(YADISK_PUBLIC_KEY, mp):
            if it.get("type") == "file": files.append(it)
    for it in files:
        n = it.get("name", "")
        if is_docx(n) and has_date(n) and not is3k(n): return it.get("path"), n
    if is3:
        for it in files:
            n = it.get("name", "")
            if is_xlsx(n) and has_date(n) and is3k(n): return it.get("path"), n
    for it in files:
        n = it.get("name", "")
        if is_docx(n) and has_date(n): return it.get("path"), n
    if is3:
        for it in files:
            n = it.get("name", "")
            if is_xlsx(n) and has_date(n): return it.get("path"), n
    return None, None


def parse_replacement_docx(content):
    out = {}
    try: doc = Document(io.BytesIO(content))
    except Exception as e:
        log.exception(f"docx: {e}"); return out
    roman = {"I":1,"II":2,"III":3,"IV":4,"V":5,"VI":6,"VII":7}
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t: continue
        lo = t.lower()
        if not ("тм 202" in lo or "тм-202" in lo or "тм202" in lo): continue
        pn = None
        for w in t.split():
            ww = w.strip(".,;:()")
            if ww in roman: pn = roman[ww]; break
        if pn is None: continue
        if "→" in t: r = t.split("→",1)[1].strip()
        elif "->" in t: r = t.split("->",1)[1].strip()
        else: r = t
        if r.lower().startswith("на "): r = r[3:].strip()
        out[pn] = r
    return out


def parse_replacement_xlsx(content):
    out = {}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        log.exception(f"xlsx: {e}"); return out
    roman = {"I":1,"II":2,"III":3,"IV":4,"V":5,"VI":6,"VII":7}
    for sh in wb.worksheets:
        for row in sh.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            line = " | ".join([c for c in cells if c.strip()])
            if not line: continue
            lo = line.lower()
            if not ("тм 202" in lo or "тм-202" in lo or "тм202" in lo): continue
            pn = None
            for w in line.split():
                ww = w.strip(".,;:()|")
                if ww in roman: pn = roman[ww]; break
            if pn is None: continue
            if "→" in line: r = line.split("→",1)[1].strip()
            elif "->" in line: r = line.split("->",1)[1].strip()
            else:
                ne = [c.strip() for c in cells if c.strip()]
                r = " | ".join(ne[-2:]) if len(ne) >= 2 else line
            r = r.strip(" |")
            if r.lower().startswith("на "): r = r[3:].strip()
            out[pn] = r
    return out


def get_replacements_for_date(target_date):
    path, name = find_replacement_file(target_date)
    if not path: return {}
    content = yadisk_download(YADISK_PUBLIC_KEY, path)
    if not content: return {}
    if name.lower().endswith(".docx"): return parse_replacement_docx(content)
    if name.lower().endswith(".xlsx"): return parse_replacement_xlsx(content)
    return {}


# ================= РАСПИСАНИЕ НА ДАТУ =================
def get_schedule_for_date(d, use_replacements=True, use_cancelled=True):
    wd = WEEKDAYS_RU[d.weekday()]
    num = is_numerator(d)
    wt = "числитель" if num else "знаменатель"
    if wd not in DAYS or d.weekday() >= 5:
        return None, wt, []
    res = []
    for it in DAYS[wd]:
        pn = it["pair"]
        if "subject" in it:
            s = it["subject"]; r = it.get("room","")
        else:
            v = it["numerator"] if num else it["denominator"]
            s = v["subject"]; r = v.get("room","")
        if s == "—": continue
        res.append({"pair": pn, "subject": s, "room": r, "time": TIMES.get(str(pn),"")})
    changes = []
    if use_cancelled:
        c = load_cancelled().get(d.strftime("%Y-%m-%d"), [])
        if c:
            rem = [p for p in res if p["pair"] in c]
            res = [p for p in res if p["pair"] not in c]
            for r in rem:
                changes.append({"pair": r["pair"], "old": r["subject"] + (f" ({r['room']})" if r["room"] else ""), "new": "ОТМЕНЕНО"})
    if use_replacements and d in (date.today(), date.today() + timedelta(days=1)):
        try:
            repl = get_replacements_for_date(d)
            if repl:
                for p in res:
                    if p["pair"] in repl:
                        old = p["subject"]
                        if p["room"]: old += f" ({p['room']})"
                        p["subject"] = repl[p["pair"]]; p["room"] = ""
                        changes.append({"pair": p["pair"], "old": old, "new": repl[p["pair"]]})
        except Exception as e:
            log.exception(f"замены: {e}")
    return res, wt, changes


# ================= ФОРМАТ =================
def format_message(d, prefix="📅 Расписание на"):
    wd = WEEKDAYS_RU[d.weekday()]
    pairs, wt, changes = get_schedule_for_date(d)
    ds = d.strftime("%d.%m.%Y")
    h = f"{prefix} {ds}, {wd.lower()} ({wt})\n\n"
    if not pairs: return h + "Пар нет 🎉"
    lines = []
    for p in pairs:
        rp = f" — {p['room']}" if p["room"] else ""
        tp = f" ({p['time']})" if p["time"] else ""
        lines.append(f"{p['pair']}. {p['subject']}{rp}{tp}")
    t = h + "\n".join(lines)
    if changes:
        t += "\n\n🔄 Изменения:\n"
        for c in changes: t += f"• {c['pair']} пара: {c['old']} → {c['new']}\n"
    t += f"\n💬 Вопросы — @{CONTACT_USERNAME}"
    return t


def format_preview(d, title):
    pairs, wt, _ = get_schedule_for_date(d, use_replacements=False)
    ds = d.strftime("%d.%m.%Y"); wd = WEEKDAYS_RU[d.weekday()]
    t = f"{title}\n{ds}, {wd.lower()} ({wt})\n\n"
    if not pairs: return t + "Пар нет 🎉"
    for p in pairs:
        rp = f" — {p['room']}" if p["room"] else ""
        t += f"{p['pair']}. {p['subject']}{rp}\n"
    return t


def format_week():
    today = date.today()
    mon = today - timedelta(days=today.weekday())
    t = f"📚 Расписание на неделю ({SCHEDULE['group']})\n\n"
    for i in range(5):
        d = mon + timedelta(days=i)
        wd = WEEKDAYS_RU[d.weekday()]
        pairs, _, _ = get_schedule_for_date(d, use_replacements=False, use_cancelled=False)
        t += f"▫️ {wd} ({d.strftime('%d.%m')})\n"
        if not pairs: t += "   — пар нет\n"
        else:
            for p in pairs:
                rp = f" — {p['room']}" if p["room"] else ""
                t += f"   {p['pair']}. {p['subject']}{rp}\n"
        t += "\n"
    t += f"💬 Вопросы — @{CONTACT_USERNAME}"
    return t.strip()


def format_bells():
    t = "🔔 Расписание звонков\n\n"
    for pn in sorted(TIMES.keys(), key=lambda x: int(x)):
        t += f"{pn} пара — {TIMES[pn]}\n"
    t += f"\n💬 Вопросы — @{CONTACT_USERNAME}"
    return t.strip()


def format_main_menu():
    return ("👋 Привет! Я бот с расписанием группы ТМ-25-202.\n\n"
            "Выбери, что тебе нужно, кнопками внизу 👇\n\n"
            f"💬 Вопросы и предложения — @{CONTACT_USERNAME}")


# ================= КНОПКИ =================
def main_keyboard():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📅 Сегодня", callback_data="btn_today"),
         InlineKeyboardButton(text="🌙 Завтра", callback_data="btn_tomorrow")],
        [InlineKeyboardButton(text="📆 Выбрать дату", callback_data="pick_year")],
        [InlineKeyboardButton(text="📚 Вся неделя", callback_data="btn_week"),
         InlineKeyboardButton(text="🔄 Изменения", callback_data="btn_changes")],
        [InlineKeyboardButton(text="🔔 Звонки", callback_data="btn_bells")],
        [InlineKeyboardButton(text="💬 Вопросы и предложения", url=f"https://t.me/{CONTACT_USERNAME}")],
    ])


def bells_keyboard():
    return InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="⬅️ Назад", callback_data="btn_back_main")]])


def years_keyboard():
    y = date.today().year
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=str(y), callback_data=f"pick_month_{y}"),
         InlineKeyboardButton(text=str(y+1), callback_data=f"pick_month_{y+1}")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="btn_back_main")],
    ])


def months_keyboard(year):
    rows = []; row = []
    for i, m in enumerate(MONTHS_RU_SHORT, start=1):
        row.append(InlineKeyboardButton(text=m, callback_data=f"pick_day_{year}_{i}"))
        if len(row) == 3: rows.append(row); row = []
    if row: rows.append(row)
    rows.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="pick_year")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def days_keyboard(year, month):
    cal = calendar.monthcalendar(year, month)
    rows = [[InlineKeyboardButton(text=x, callback_data="noop") for x in ["Пн","Вт","Ср","Чт","Пт","Сб","Вс"]]]
    for w in cal:
        row = []
        for d in w:
            if d == 0: row.append(InlineKeyboardButton(text=" ", callback_data="noop"))
            else: row.append(InlineKeyboardButton(text=str(d), callback_data=f"show_date_{year}_{month}_{d}"))
        rows.append(row)
    rows.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=f"pick_month_{year}")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


# ================= РАССЫЛКА =================
async def broadcast(bot, text):
    subs = load_subscribers()
    if not subs: return
    for cid in subs:
        if is_muted(cid): continue
        try: await bot.send_message(chat_id=cid, text=text, reply_markup=main_keyboard())
        except Exception as e:
            log.warning(f"{cid}: {e}")
            if "blocked" in str(e).lower() or "chat not found" in str(e).lower():
                remove_subscriber(cid)


async def notify_owner(bot, text):
    cid = get_owner_chat_id()
    if not cid: return
    try: await bot.send_message(chat_id=cid, text=text)
    except Exception as e: log.warning(f"owner: {e}")


async def send_today(bot):
    await broadcast(bot, format_message(date.today(), "☀️ Расписание на сегодня:"))
    await notify_owner(bot, f"✅ Рассылка «на сегодня». Подписчиков: {len(load_subscribers())}")


async def send_tomorrow(bot):
    await broadcast(bot, format_message(date.today() + timedelta(days=1), "🌙 Расписание на завтра:"))
    await notify_owner(bot, f"✅ Рассылка «на завтра». Подписчиков: {len(load_subscribers())}")


async def preview_today(bot):
    await notify_owner(bot, format_preview(date.today(), "🔔 Через 10 минут — рассылка на СЕГОДНЯ:"))


async def preview_tomorrow(bot):
    await notify_owner(bot, format_preview(date.today() + timedelta(days=1), "🔔 Через 10 минут — рассылка на ЗАВТРА:"))


async def warn_today(bot): await notify_owner(bot, "⏰ Через 2 минуты — рассылка «на сегодня».")
async def warn_tomorrow(bot): await notify_owner(bot, "⏰ Через 2 минуты — рассылка «на завтра».")


# ================= MIDDLEWARE =================
class BanMiddleware(BaseMiddleware):
    async def __call__(self, handler, event, data):
        if isinstance(event, Message) and is_banned(event):
            return
        return await handler(event, data)


# ================= БОТ =================
bot = Bot(token=TELEGRAM_TOKEN)
dp = Dispatcher()
dp.message.middleware(BanMiddleware())
dp.callback_query.middleware(BanMiddleware())
dp.inline_query.middleware(BanMiddleware())


@dp.message(CommandStart())
async def cmd_start(m: Message):
    add_subscriber(m.chat.id)
    if is_owner(m):
        set_owner_chat_id(m.chat.id)
        await m.answer("👑 Ты — ВЛАДЕЛЕЦ.\nКоманды: /ownerhelp")
    elif is_admin(m):
        await m.answer("🛠 Ты — админ.\nКоманды: /edithelp")
    await m.answer(format_main_menu(), reply_markup=main_keyboard())


@dp.message(Command("menu"))
async def cmd_menu(m): await m.answer(format_main_menu(), reply_markup=main_keyboard())


@dp.message(Command("stop"))
async def cmd_stop(m): remove_subscriber(m.chat.id); await m.answer("Отписался.")


@dp.message(Command("today"))
async def cmd_today(m): await m.answer(format_message(date.today()), reply_markup=main_keyboard())


@dp.message(Command("tomorrow"))
async def cmd_tomorrow(m): await m.answer(format_message(date.today()+timedelta(days=1)), reply_markup=main_keyboard())


@dp.message(Command("week_all"))
async def cmd_week_all(m): await m.answer(format_week(), reply_markup=main_keyboard())


@dp.message(Command("bells"))
async def cmd_bells(m): await m.answer(format_bells(), reply_markup=bells_keyboard())


@dp.message(Command("week"))
async def cmd_week(m):
    t = date.today(); n = is_numerator(t)
    wt = "числитель" if n else "знаменатель"
    nxt = "знаменатель" if n else "числитель"
    await m.answer(f"📆 Сейчас неделя: {wt}\nДо смены на {nxt}: {days_until_week_change(t)} дн.\nСегодня: {WEEKDAYS_RU[t.weekday()]}, {t.strftime('%d.%m.%Y')}")


@dp.message(Command("find"))
async def cmd_find(m):
    parts = m.text.split(maxsplit=1)
    if len(parts) < 2: await m.answer("Использование: /find <предмет>"); return
    q = parts[1].lower(); found = []
    for wd, pairs in DAYS.items():
        for it in pairs:
            if "subject" in it:
                if q in it["subject"].lower(): found.append((wd, it["pair"], it["subject"], it.get("room","")))
            else:
                for v in ("numerator","denominator"):
                    vv = it[v]
                    if q in vv["subject"].lower():
                        lbl = "числ" if v=="numerator" else "знам"
                        found.append((wd, it["pair"], f"{vv['subject']} [{lbl}]", vv.get("room","")))
    if not found: await m.answer(f"Ничего не найдено."); return
    t = f"🔍 Найдено:\n\n"
    for wd, pn, s, r in found:
        t += f"• {wd}, {pn} пара: {s}" + (f" — {r}" if r else "") + "\n"
    await m.answer(t)


@dp.message(Command("date"))
async def cmd_date(m):
    parts = m.text.split(maxsplit=1)
    if len(parts) < 2: await m.answer("Использование: /date <ДД.ММ>"); return
    try:
        d, mo = parts[1].split(".")
        t = date(date.today().year, int(mo), int(d))
    except Exception: await m.answer("Формат: ДД.ММ"); return
    await m.answer(format_message(t), reply_markup=main_keyboard())


@dp.message(Command("mute"))
async def cmd_mute(m):
    parts = m.text.split()
    if len(parts) < 2: await m.answer("/mute <дни>"); return
    try: days = int(parts[1])
    except: await m.answer("Число."); return
    until = (date.today()+timedelta(days=days)).strftime("%Y-%m-%d")
    mm = load_mute(); mm[str(m.chat.id)] = until; save_mute(mm)
    await m.answer(f"🔕 Тишина до {until}")


@dp.message(Command("unmute"))
async def cmd_unmute(m):
    mm = load_mute(); mm.pop(str(m.chat.id), None); save_mute(mm)
    await m.answer("🔔 Включено.")


@dp.message(Command("ownerhelp"))
async def cmd_ownerhelp(m):
    if not is_owner(m): return
    await m.answer("/ping /logs /errors /stats /addadmin /deladmin /listadmins /ban /banid /unban /listbanned /broadcast\nВсе админ-команды: /edithelp")


@dp.message(Command("ping"))
async def cmd_ping(m):
    if not is_owner(m): return
    up = int(time.time()-START_TIME)
    await m.answer(f"🏓 Понг! Аптайм: {up//3600}ч {(up%3600)//60}м {up%60}с\nПодписчиков: {len(load_subscribers())}")


@dp.message(Command("stats"))
async def cmd_stats(m):
    if not is_owner(m): return
    s = load_settings()
    await m.answer(f"📊 Подписчиков: {len(load_subscribers())}\nАдминов: {len(load_admins())}\nБан: {len(load_banned())+len(load_banned_usernames())}\nТочка: {s['ref_monday']}")


def _norm_u(raw): return raw.strip().lstrip("@").lower()


@dp.message(Command("addadmin"))
async def cmd_addadmin(m):
    if not is_owner(m): return
    p = m.text.split(maxsplit=1)
    if len(p) < 2: await m.answer("/addadmin @username"); return
    u = _norm_u(p[1]); a = load_admins()
    if u in a: await m.answer("Уже админ."); return
    a.add(u); save_admins(a); await m.answer(f"✅ @{u}")


@dp.message(Command("deladmin"))
async def cmd_deladmin(m):
    if not is_owner(m): return
    p = m.text.split(maxsplit=1)
    if len(p) < 2: await m.answer("/deladmin @username"); return
    u = _norm_u(p[1]); a = load_admins()
    if u not in a: await m.answer("Не админ."); return
    a.remove(u); save_admins(a); await m.answer(f"🗑 @{u}")


@dp.message(Command("listadmins"))
async def cmd_listadmins(m):
    if not is_owner(m): return
    a = sorted(load_admins())
    await m.answer("🛠 " + (", ".join("@"+x for x in a) if a else "нет"))


@dp.message(Command("ban"))
async def cmd_ban(m):
    if not is_owner(m): return
    p = m.text.split(maxsplit=1)
    if len(p) < 2: await m.answer("/ban @username"); return
    u = _norm_u(p[1]); bu = load_banned_usernames(); bu.add(u); save_banned_usernames(bu)
    await m.answer(f"🚫 @{u}")


@dp.message(Command("banid"))
async def cmd_banid(m):
    if not is_owner(m): return
    p = m.text.split()
    if len(p) < 2: await m.answer("/banid <id>"); return
    try: cid = int(p[1])
    except: await m.answer("Число."); return
    b = load_banned(); b.add(cid); save_banned(b); remove_subscriber(cid)
    await m.answer(f"🚫 {cid}")


@dp.message(Command("unban"))
async def cmd_unban(m):
    if not is_owner(m): return
    p = m.text.split(maxsplit=1)
    if len(p) < 2: await m.answer("/unban @username"); return
    u = _norm_u(p[1]); bu = load_banned_usernames()
    if u in bu: bu.remove(u); save_banned_usernames(bu); await m.answer(f"✅ @{u}")
    else: await m.answer("Не в бан-листе.")


@dp.message(Command("listbanned"))
async def cmd_listbanned(m):
    if not is_owner(m): return
    await m.answer("🚫 " + (", ".join(sorted(load_banned_usernames())) or "пусто"))


@dp.message(Command("broadcast"))
async def cmd_broadcast(m):
    if not is_owner(m): return
    p = m.text.split(maxsplit=1)
    if len(p) < 2: await m.answer("/broadcast <текст>"); return
    sent = 0
    for cid in load_subscribers():
        try: await bot.send_message(cid, p[1]); sent += 1
        except: pass
    await m.answer(f"📣 Отправлено: {sent}")


@dp.message(Command("edithelp"))
async def cmd_edithelp(m):
    if not is_admin(m): return
    await m.answer("/setpair /delpair /settime /setweek /show /replacements /cancel /uncancel /setref")


def _norm_day(raw): return WEEKDAYS_ALIASES.get(raw.strip().lower())


@dp.message(Command("show"))
async def cmd_show(m):
    if not is_admin(m): return
    p = m.text.split(maxsplit=1)
    if len(p) < 2: await m.answer("/show <день>"); return
    d = _norm_day(p[1])
    if not d: await m.answer("Не понял."); return
    pairs = DAYS.get(d, [])
    if not pairs: await m.answer(f"{d}: нет"); return
    t = f"📋 {d}:\n\n"
    for it in pairs:
        pn = it["pair"]; tt = TIMES.get(str(pn),"")
        if "subject" in it: t += f"{pn}. {it['subject']} — {it.get('room','')} ({tt})\n"
        else:
            t += f"{pn}. [числ] {it['numerator']['subject']}\n"
            t += f"   [знам] {it['denominator']['subject']}\n"
    await m.answer(t)


@dp.message(Command("setpair"))
async def cmd_setpair(m):
    if not is_admin(m): return
    p = m.text.split(maxsplit=4)
    if len(p) < 5: await m.answer("/setpair <день> <пара> <предмет> <аудитория>"); return
    _, dr, pr, subj, room = p
    d = _norm_day(dr)
    if not d: await m.answer("Не понял день."); return
    try: pn = int(pr)
    except: await m.answer("Число."); return
    dl = DAYS.setdefault(d, [])
    ok = False
    for it in dl:
        if it["pair"] == pn:
            it.clear(); it.update({"pair": pn, "subject": subj, "room": room}); ok = True; break
    if not ok:
        dl.append({"pair": pn, "subject": subj, "room": room}); dl.sort(key=lambda x: x["pair"])
    await m.answer(f"✅ {d}, {pn} → {subj}, {room}")


@dp.message(Command("delpair"))
async def cmd_delpair(m):
    if not is_admin(m): return
    p = m.text.split()
    if len(p) < 3: await m.answer("/delpair <день> <пара>"); return
    d = _norm_day(p[1])
    if not d: await m.answer("Не понял."); return
    try: pn = int(p[2])
    except: await m.answer("Число."); return
    dl = DAYS.get(d, []); before = len(dl)
    DAYS[d] = [x for x in dl if x["pair"] != pn]
    if len(DAYS[d]) == before: await m.answer("Не найдено."); return
    await m.answer(f"🗑 {d}, {pn}")


@dp.message(Command("settime"))
async def cmd_settime(m):
    if not is_admin(m): return
    p = m.text.split(maxsplit=2)
    if len(p) < 3: await m.answer("/settime <пара> <время>"); return
    try: pn = int(p[1])
    except: await m.answer("Число."); return
    TIMES[str(pn)] = p[2].strip()
    await m.answer(f"✅ {pn} → {p[2]}")


@dp.message(Command("setweek"))
async def cmd_setweek(m):
    if not is_admin(m): return
    p = m.text.split(maxsplit=5)
    if len(p) < 6: await m.answer("/setweek <день> <пара> <numerator|denominator> <предмет> <аудитория>"); return
    _, dr, pr, vr, subj, room = p
    d = _norm_day(dr)
    if not d: await m.answer("Не понял."); return
    if vr not in ("numerator","denominator"): await m.answer("numerator|denominator"); return
    try: pn = int(pr)
    except: await m.answer("Число."); return
    dl = DAYS.setdefault(d, [])
    ok = False
    for it in dl:
        if it["pair"] == pn and "numerator" in it:
            it[vr] = {"subject": subj, "room": room}; ok = True; break
    if not ok:
        new = {"pair": pn, "numerator": {"subject":"—","room":""}, "denominator": {"subject":"—","room":""}}
        new[vr] = {"subject": subj, "room": room}
        DAYS[d] = [x for x in dl if x["pair"] != pn] + [new]
        DAYS[d].sort(key=lambda x: x["pair"])
    await m.answer(f"✅ {d}, {pn} ({vr}) → {subj}, {room}")


@dp.message(Command("cancel"))
async def cmd_cancel(m):
    if not is_admin(m): return
    p = m.text.split()
    if len(p) < 3: await m.answer("/cancel <ДД.ММ> <пара>"); return
    try:
        d, mo = p[1].split("."); t = date(date.today().year, int(mo), int(d))
        pn = int(p[2])
    except: await m.answer("Формат: ДД.ММ и число."); return
    key = t.strftime("%Y-%m-%d"); c = load_cancelled()
    c.setdefault(key, [])
    if pn not in c[key]: c[key].append(pn)
    save_cancelled(c); await m.answer(f"🚫 {key}, {pn}")


@dp.message(Command("uncancel"))
async def cmd_uncancel(m):
    if not is_admin(m): return
    p = m.text.split()
    if len(p) < 3: await m.answer("/uncancel <ДД.ММ> <пара>"); return
    try:
        d, mo = p[1].split("."); t = date(date.today().year, int(mo), int(d))
        pn = int(p[2])
    except: await m.answer("Формат: ДД.ММ и число."); return
    key = t.strftime("%Y-%m-%d"); c = load_cancelled()
    if key in c and pn in c[key]:
        c[key].remove(pn)
        if not c[key]: c.pop(key)
        save_cancelled(c); await m.answer(f"✅ {key}, {pn}")
    else: await m.answer("Не найдено.")


@dp.message(Command("setref"))
async def cmd_setref(m):
    if not is_admin(m): return
    p = m.text.split()
    if len(p) < 3: await m.answer("/setref <ДД.ММ.ГГГГ> <числитель|знаменатель>"); return
    try: rd = datetime.strptime(p[1], "%d.%m.%Y").date()
    except: await m.answer("ДД.ММ.ГГГГ"); return
    v = p[2].lower()
    if v not in ("числитель","знаменатель"): await m.answer("числитель|знаменатель"); return
    s = load_settings(); s["ref_monday"] = rd.strftime("%Y-%m-%d"); s["ref_is_numerator"] = (v=="числитель"); save_settings(s)
    await m.answer(f"✅ {rd} — {v}")


@dp.message(Command("replacements"))
async def cmd_replacements(m):
    if not is_admin(m): return
    await m.answer("⏳ Ищу...")
    path, name = find_replacement_file(date.today())
    if not path: await m.answer("❌ Файл не найден."); return
    content = yadisk_download(YADISK_PUBLIC_KEY, path)
    if not content: await m.answer("❌ Не скачалось."); return
    if name.lower().endswith(".docx"): r = parse_replacement_docx(content)
    elif name.lower().endswith(".xlsx"): r = parse_replacement_xlsx(content)
    else: r = {}
    if not r: await m.answer(f"📄 {name}\n\nЗамен нет."); return
    t = f"📄 {name}\n\n"
    for pn, x in sorted(r.items()): t += f"• {pn} → {x}\n"
    await m.answer(t)


@dp.inline_query()
async def inline_handler(q: InlineQuery):
    qs = q.query.strip().lower(); today = date.today(); tm = today + timedelta(days=1)
    res = []
    if qs in ("","сегодня","today"):
        res.append(InlineQueryResultArticle(id="t", title="Сегодня", input_message_content=InputTextMessageContent(message_text=format_message(today))))
    if qs in ("","завтра","tomorrow"):
        res.append(InlineQueryResultArticle(id="tm", title="Завтра", input_message_content=InputTextMessageContent(message_text=format_message(tm))))
    if qs in ("","неделя","week"):
        res.append(InlineQueryResultArticle(id="w", title="Неделя", input_message_content=InputTextMessageContent(message_text=format_week())))
    await q.answer(res, cache_time=60)


@dp.callback_query()
async def cb(c: CallbackQuery):
    a = c.data
    text = format_main_menu(); markup = main_keyboard()
    if a == "btn_today": text = format_message(date.today())
    elif a == "btn_tomorrow": text = format_message(date.today()+timedelta(days=1))
    elif a == "btn_week": text = format_week()
    elif a == "btn_bells": text = format_bells(); markup = bells_keyboard()
    elif a == "btn_changes":
        _, _, ch = get_schedule_for_date(date.today())
        if not ch: text = "🔄 Изменений нет."
        else:
            text = "🔄 Изменения:\n\n"
            for x in ch: text += f"• {x['pair']} пара: {x['old']} → {x['new']}\n"
        text += f"\n💬 @{CONTACT_USERNAME}"
    elif a == "btn_back_main": text = format_main_menu()
    elif a == "pick_year": text = "Укажи год:"; markup = years_keyboard()
    elif a.startswith("pick_month_"):
        y = int(a.split("_")[2]); text = f"Месяц ({y}):"; markup = months_keyboard(y)
    elif a.startswith("pick_day_"):
        _, _, y, mo = a.split("_"); y, mo = int(y), int(mo)
        text = f"День ({MONTHS_RU_SHORT[mo-1]} {y}):"; markup = days_keyboard(y, mo)
    elif a.startswith("show_date_"):
        _, _, y, mo, d = a.split("_")
        text = format_message(date(int(y), int(mo), int(d)))
    elif a == "noop":
        await c.answer(); return
    try:
        await c.message.edit_text(text, reply_markup=markup)
    except Exception as e:
        if "message is not modified" not in str(e).lower():
            try: await c.message.answer(text, reply_markup=markup)
            except: pass
    await c.answer()


# ================= ПЛАНИРОВЩИК =================
scheduler = None


def init_scheduler():
    global scheduler
    scheduler = AsyncIOScheduler(timezone="Asia/Yekaterinburg")
    scheduler.add_job(preview_today, "cron", hour=5, minute=50, args=[bot])
    scheduler.add_job(preview_tomorrow, "cron", hour=17, minute=50, args=[bot])
    scheduler.add_job(warn_today, "cron", hour=5, minute=58, args=[bot])
    scheduler.add_job(warn_tomorrow, "cron", hour=17, minute=58, args=[bot])
    scheduler.add_job(send_today, "cron", hour=6, minute=0, args=[bot])
    scheduler.add_job(send_tomorrow, "cron", hour=18, minute=0, args=[bot])
    scheduler.start()
    log.info("Планировщик запущен.")
